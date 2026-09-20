"""仅修复渲染副本的文字自适应；不改文字、公式对象、电路或图像内容。"""
import sys,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'python_libs'))
from pptx import Presentation
from pptx.util import Pt
from PIL import ImageFont
from pptx.enum.text import MSO_AUTO_SIZE
def repair(src,dest):
    prs=Presentation(src);changes=[]
    for n,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if not sh.has_text_frame or not sh.text.strip():continue
            tf=sh.text_frame
            # 以原始文本的显式换行为准，保留所有上下标与字符。
            w=(sh.width-tf.margin_left-tf.margin_right)/12700
            h=(sh.height-tf.margin_top-tf.margin_bottom)/12700
            if w<=0 or h<=0:continue
            maxwidth=0;height=0;sizes=[]
            for p in tf.paragraphs:
                line_width=0;maxsize=0
                for r in p.runs:
                    sz=r.font.size.pt if r.font.size else (p.font.size.pt if p.font.size else 24)
                    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',int(sz*4))
                    line_width+=font.getlength(r.text)/4;maxsize=max(maxsize,sz);sizes.append((r,sz))
                # 显式换行需要多行预算。以略保守的中文字体测量避免替代字体外溢。
                count=max(1,p.text.count('\v')+1)
                maxwidth=max(maxwidth,line_width/count)
                height+=maxsize*1.25*count+(p.space_after.pt if p.space_after else 0)+(p.space_before.pt if p.space_before else 0)
            scale=min(1,w/maxwidth if maxwidth else 1,h/height if height else 1)
            if scale<.97:
                tf.auto_size=MSO_AUTO_SIZE.NONE
                for r,sz in sizes:r.font.size=Pt(max(9,sz*scale*.94))
                changes.append({'page':n,'shape':sh.name,'scale':round(scale*.94,3),'text':sh.text})
    prs.save(dest)
    Path(str(dest)+'.changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),'utf8')
    print(dest,len(changes),'text frames fitted')
if __name__=='__main__':repair(sys.argv[1],sys.argv[2])
