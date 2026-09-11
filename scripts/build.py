#!/usr/bin/env python3
"""按 outline.json 生成「通俗讲解版」PPTX。

outline.json 结构:
{
  "title": "...",
  "subtitle": "...",
  "output": "xxx(通俗版).pptx",
  "slides": [
    {"type": "cover"},
    {"type": "concept", "title": "...", "one_liner": "一句话版本",
     "bullets": ["..."], "analogy": "生活类比/人话解释"},
    {"type": "formula", "title": "...",
     "source": "assets/pages/slide_012.png",   # 优先:原页截图,无损
     "latex": "x(t)=\\sum...",                 # 备选:matplotlib mathtext 渲染
     "note": "每个符号的人话解释"},
    {"type": "example", "title": "例题(原第12页)",
     "source": "assets/pages/slide_012.png",
     "source_label": "原课件第12页",
     "analysis": "这题在考什么 / 思路",
     "steps": ["步骤1", "步骤2"]},
    {"type": "summary", "title": "本章小结", "bullets": ["..."]}
  ]
}

用法:
    python build.py --outline outline.json --work ./work
输出:
    ./work/<output>
"""
import argparse
import json
import math
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = 13.333, 7.5

NAVY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xFB, 0xFC, 0xFD)
NOTE_BG = RGBColor(0xFF, 0xF7, 0xE0)   # 淡黄,类比/人话框
NOTE_LINE = RGBColor(0xB9, 0x90, 0x00)
TERM_BG = RGBColor(0xEA, 0xF1, 0xF8)   # 浅蓝,名词小抄框
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
EXAM_STYLE = {                          # 重点标签: (底色, 字色)
    "必考": (RGBColor(0xC0, 0x39, 0x2B), WHITE),
    "掌握": (RGBColor(0xFF, 0xC0, 0x00), RGBColor(0x7F, 0x30, 0x00)),
    "了解": (RGBColor(0xD6, 0xD6, 0xD6), DARK),
}


def add_header(slide, title, page_no=None, exam=None):
    """深蓝标题条 + 标题文字 + 页码 + 可选重点标签(必考/掌握/了解)。"""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                 Inches(SLIDE_W), Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.16), Inches(10.4), Inches(0.75))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(20 if len(title) > 34 else 22 if len(title) > 25
                     else 26 if len(title) <= 18 else 24)
    r.font.bold = True
    r.font.color.rgb = WHITE
    if exam:
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(11.2), Inches(0.27), Inches(1.0), Inches(0.52))
        tag.fill.solid()
        tag.fill.fore_color.rgb = EXAM_STYLE.get(exam, EXAM_STYLE["了解"])[0]
        tag.line.fill.background()
        tag_tf = tag.text_frame
        tag_tf.margin_left = Inches(0.05)
        tag_tf.margin_right = Inches(0.05)
        tp = tag_tf.paragraphs[0]
        tp.alignment = PP_ALIGN.CENTER
        tr = tp.add_run()
        tr.text = exam
        tr.font.size = Pt(14)
        tr.font.bold = True
        tr.font.color.rgb = EXAM_STYLE.get(exam, EXAM_STYLE["了解"])[1]
    if page_no:
        pb = slide.shapes.add_textbox(Inches(12.3), Inches(0.25), Inches(0.9), Inches(0.55))
        pp = pb.text_frame.paragraphs[0]
        pp.alignment = PP_ALIGN.RIGHT
        rp = pp.add_run()
        rp.text = str(page_no)
        rp.font.size = Pt(11)
        rp.font.color.rgb = WHITE


def fit_pt(texts, width_in, height_in, base_pt):
    """根据文本量与容器估算合适字号(pptx 无法自动排版,保证不溢出)。"""
    for pt in range(base_pt, base_pt - 13, -2):
        chars_per_line = max(4.0, width_in * 72 / (pt * 1.02))
        lines = sum(max(1, math.ceil(len(t) / chars_per_line)) for t in texts)
        if lines * pt * 1.32 / 72 <= height_in:
            return pt
    return base_pt - 12


def fill_bullets(tf, texts, pt, color=DARK, bullet="· "):
    first = True
    for t in texts:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = bullet + t
        r.font.size = Pt(pt)
        r.font.color.rgb = color


def add_note(slide, text, y, h=0.78, label="💡 类比"):
    """底部浅黄色「人话」条。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(0.55), Inches(y), Inches(12.23), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NOTE_BG
    shape.line.color.rgb = NOTE_LINE
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"{label}: {text}"
    r.font.size = Pt(fit_pt([f"{label}: {text}"], 11.9, h, 16))
    r.font.color.rgb = DARK


def add_terms(slide, terms, y):
    """底部浅蓝色「📚 名词小抄」框:一条一个名词:解释。"""
    h = 0.34 + 0.31 * len(terms)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(0.55), Inches(y), Inches(12.23), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TERM_BG
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "📚 名词小抄"
    r0.font.size = Pt(13)
    r0.font.bold = True
    r0.font.color.rgb = NAVY
    pt = fit_pt(terms, 11.9, h - 0.4, 12)
    for t in terms:
        p = tf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = "· " + t
        r.font.size = Pt(pt)
        r.font.color.rgb = DARK


def picture_fit(slide, src, left, top, max_w, max_h):
    """等比缩放插入图片,使其落在 (left, top) 最多 max_w x max_h 的盒子里。"""
    with Image.open(src) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    left = left + (max_w - w) / 2
    top = top + (max_h - h) / 2
    return slide.shapes.add_picture(str(src), Inches(left), Inches(top),
                                    Inches(w), Inches(h)), (w, h)


def latex_to_png(latex, out_png):
    """用 matplotlib mathtext 渲染 LaTeX 公式为透明 PNG。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${latex}$", fontsize=32)
    fig.savefig(out_png, dpi=200, bbox_inches="tight", pad_inches=0.15, transparent=True)
    plt.close(fig)


def build_slide(prs, slide_def, work, page_no):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kind = slide_def["type"]

    if kind == "cover":
        t = slide_def.get("title", "通俗版课件")
        sub = slide_def.get("subtitle", "")
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.6))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = t
        r.font.size = Pt(40)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if sub:
            sb = slide.shapes.add_textbox(Inches(1.5), Inches(4.0), Inches(10.3), Inches(1.2))
            stf = sb.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.alignment = PP_ALIGN.CENTER
            sr = sp.add_run()
            sr.text = sub
            sr.font.size = Pt(18)
            sr.font.color.rgb = GRAY
        return

    add_header(slide, slide_def.get("title", ""), page_no,
               exam=slide_def.get("exam"))
    body_top = 1.28
    body_h = 5.05 if kind in ("concept", "summary") else None

    if kind in ("concept", "summary"):
        terms = slide_def.get("terms") or []
        note = slide_def.get("analogy") or slide_def.get("note")
        has_note = bool(note) and kind == "concept"
        top = slide_def.get("one_liner")
        body_top = 1.28
        if top:
            ob = slide.shapes.add_textbox(Inches(0.6), Inches(1.18), Inches(12.1), Inches(0.4))
            otf = ob.text_frame
            otf.word_wrap = True
            op = otf.paragraphs[0]
            or_ = op.add_run()
            or_.text = f"一句话: {top}"
            or_.font.size = Pt(15)
            or_.font.italic = True
            or_.font.color.rgb = GRAY
            body_top = 1.62
        # 垂直布局预算:自下而上给「名词小抄」「类比条」留位
        if terms:
            terms_h = 0.34 + 0.31 * len(terms)
            terms_y = max(body_top + 1.7, 7.42 - terms_h)
            note_h = 0.62 if has_note else 0
            note_y = (terms_y - 0.08 - note_h) if has_note else None
            body_bottom = (note_y if has_note else terms_y) - 0.1
        else:
            note_h = 0.78 if has_note else 0
            note_y = 6.5 if has_note else None
            body_bottom = (note_y - 0.05 if has_note else 7.4)
        body_h = max(1.5, body_bottom - body_top)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(body_top),
                                       Inches(12.1), Inches(body_h))
        tf = box.text_frame
        tf.word_wrap = True
        pt = fit_pt(slide_def.get("bullets", [""]), 12.1, body_h, 20)
        fill_bullets(tf, slide_def.get("bullets", []), pt)
        if has_note:
            add_note(slide, note, note_y, h=note_h)
        if terms:
            add_terms(slide, terms, terms_y)

    elif kind == "formula":
        src = work / slide_def.get("source", "") if slide_def.get("source") else None
        if src and Path(src).exists():
            picture_fit(slide, src, 0.8, 1.35, 11.7, 4.4)
        elif slide_def.get("latex"):
            png = work / "assets" / f"formula_{page_no}.png"
            try:
                latex_to_png(slide_def["latex"], str(png))
            except Exception as e:
                sys.exit(f"formula 渲染失败(需 pip install matplotlib): {e}")
            picture_fit(slide, png, 0.8, 1.35, 11.7, 4.4)
        note = slide_def.get("note")
        if note:
            add_note(slide, note, 6.25, label="📖 符号解释")

    elif kind == "example":
        src = slide_def.get("source")
        sources = [src] if isinstance(src, str) else (src or [])
        src_paths = [work / s for s in sources]
        if src_paths and Path(src_paths[0]).exists():
            n = len(src_paths)
            pic_h = 4.55
            box_w = (11.7 - 0.15 * (n - 1)) / n
            for k, sp in enumerate(src_paths):
                if not Path(sp).exists():
                    print(f"warn: 例题原图缺失 {sp}", file=sys.stderr)
                    continue
                picture_fit(slide, sp, 0.8 + k * (box_w + 0.15), 1.3, box_w, pic_h)
        label = slide_def.get("source_label", "")
        if label:
            lb = slide.shapes.add_textbox(Inches(0.6), Inches(1.06), Inches(12.1), Inches(0.28))
            ltf = lb.text_frame
            lp = ltf.paragraphs[0]
            lr = lp.add_run()
            lr.text = f"📌 {label}"
            lr.font.size = Pt(12)
            lr.font.color.rgb = GRAY
        lines = [f"解析: {slide_def.get('analysis', '')}"] if slide_def.get("analysis") else []
        for i, s in enumerate(slide_def.get("steps", []), 1):
            lines.append(f"  第{i}步: {s}")
        if lines:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(6.0), Inches(12.1), Inches(1.1))
            tf = box.text_frame
            tf.word_wrap = True
            pt = fit_pt(lines, 12.1, 1.1, 15)
            fill_bullets(tf, lines, pt, bullet="")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--outline", required=True, help="outline.json 路径")
    ap.add_argument("--work", required=True, help="工作目录(与 extract.py 相同)")
    args = ap.parse_args()

    outline = json.loads(Path(args.outline).read_text(encoding="utf-8"))
    work = Path(args.work)
    out_name = outline.get("output", "通俗版.pptx")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    total = len(outline.get("slides", []))
    for i, s in enumerate(outline.get("slides", []), 1):
        build_slide(prs, s, work, i)
    out_path = work / out_name
    prs.save(str(out_path))

    print(f"已生成 {total} 页 -> {out_path}")

    # 顺手渲染成 PNG,供下一步检查(复用 extract 的渲染脚本)
    try:
        subprocess.run([sys.executable, str(Path(__file__).parent / "render.py"),
                        "--input", str(out_path), "--out", str(work / "check")],
                       check=False)
    except Exception:
        pass

    # 几何布局检查:文本框溢出 / 越界 / 重叠
    check = Path(__file__).parent / "check_layout.py"
    r = subprocess.run([sys.executable, str(check), "--input", str(out_path)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        print("⚠ ", file=sys.stderr)
        print((r.stdout or r.stderr) or "布局检查发现问题", file=sys.stderr)
        print("提示:检查上面的溢出/重叠报告,修改 outline.json(通常缩短文字)后重跑 build。", file=sys.stderr)
        sys.exit(1)
    print(r.stdout.strip())


if __name__ == "__main__":
    main()
