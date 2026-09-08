#!/usr/bin/env python3
"""几何布局检查(不依赖视觉):文本框溢出、图片越界、元素重叠。

与 build.py 的音量一致:字宽≈1.02pt、行高≈1.32pt 估算,超过容器则报告。
用法:
    python check_layout.py --input out.pptx
退出码 0 = 通过;1 = 有问题。
"""
import argparse
import math
import sys
from pathlib import Path

from pptx import Presentation

EMU_IN = 914400


def est_lines(text: str, width_in: float, pt: float) -> int:
    chars_per_line = max(4.0, width_in * 72 / (pt * 1.02))
    return max(1, math.ceil(len(text) / chars_per_line))


def est_height_in(text: str, width_in: float, pt: float) -> float:
    return est_lines(text, width_in, pt) * pt * 1.32 / 72


def runs_pt(shape) -> float:
    pts = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                pts.append(r.font.size.pt)
    return max(pts) if pts else 18.0


def in_(emu) -> float:
    return emu / EMU_IN


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True)
    args = ap.parse_args()

    prs = Presentation(args.input)
    sw, slide_h = in_(prs.slide_width), in_(prs.slide_height)
    problems = []

    for si, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        for shape_ in shapes:
            left, top = in_(shape_.left), in_(shape_.top)
            w, h = in_(shape_.width), in_(shape_.height)
            if left < -0.01 or top < -0.01 or left + w > sw + 0.01 or top + h > slide_h + 0.01:
                problems.append(f"第{si}页:形状越界 left={left:.2f} top={top:.2f} w={w:.2f} h={h:.2f}")

            # 文本框溢出
            if shape_.has_text_frame and shape_.text_frame.text.strip():
                pt = runs_pt(shape_)
                avail_w = max(0.2, w - 0.2)  # 默认左右边距
                need = est_height_in(shape_.text_frame.text, avail_w, pt)
                if need > h + 0.03:
                    problems.append(
                        f"第{si}页:文本框可能溢出 '{shape_.text_frame.text[:24]}…' "
                        f"需 {need:.2f}in 有 {h:.2f}in (字号{pt:.0f})")

        # 重叠:内容区文本框之间/与图片之间
        for i in range(len(shapes)):
            for j in range(i + 1, len(shapes)):
                a, b = shapes[i], shapes[j]
                if not (a.has_text_frame and a.text_frame.text.strip()):
                    continue
                if not (b.has_text_frame or (b.shape_type == 13)):  # 13=图片
                    continue
                ax0, ay0, ax1, ay1 = in_(a.left), in_(a.top), in_(a.left + a.width), in_(a.top + a.height)
                bx0, by0, bx1, by1 = in_(b.left), in_(b.top), in_(b.left + b.width), in_(b.top + b.height)
                ox = max(0, min(ax1, bx1) - max(ax0, bx0))
                oy = max(0, min(ay1, by1) - max(ay0, by0))
                if ox > 0.05 and oy > 0.05:
                    problems.append(
                        f"第{si}页:文本框与{'文本框' if b.has_text_frame else '图片'}重叠 "
                        f"{ox:.2f}x{oy:.2f}in '{a.text_frame.text[:16]}…'")

    if problems:
        print("布局问题:")
        for p in problems:
            print(" -", p)
        sys.exit(1)
    print("布局检查通过:无溢出/越界/重叠。")


if __name__ == "__main__":
    main()
